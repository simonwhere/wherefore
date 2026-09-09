#!/usr/bin/env python3
"""
audit_workbook.py
-----------------
Data-integrity audit for SOPO_review_raw_analysis.xlsx.

Checks the workbook against itself (cross-sheet consistency), against the
provenance rules stated in build_sopo_workbook.py's docstring, and against
facts that can be independently derived (TikTok video IDs encode a post
timestamp).

Run:  python3 audit_workbook.py [path/to/workbook.xlsx]
Exit code 1 if any ERROR-level finding is present.
"""

import datetime as dt
import re
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

PATH = sys.argv[1] if len(sys.argv) > 1 else "SOPO_review_raw_analysis.xlsx"

ERRORS, WARNS, INFOS = [], [], []


def err(code, msg):
    ERRORS.append((code, msg))


def warn(code, msg):
    WARNS.append((code, msg))


def info(code, msg):
    INFOS.append((code, msg))


def sheet_records(ws):
    """Return list of dicts keyed by header row."""
    rows = list(ws.iter_rows(values_only=True))
    headers = [h for h in rows[0]]
    out = []
    for r in rows[1:]:
        out.append({h: (v if v is not None else "") for h, v in zip(headers, r)})
    return headers, out


wb = load_workbook(PATH)

# --------------------------------------------------------------------------
# 0. Structure
# --------------------------------------------------------------------------
EXPECTED_SHEETS = ["raw_reviews", "cleaned_reviews", "theme_tags", "source_log"]
if wb.sheetnames != EXPECTED_SHEETS:
    err("S1", f"Sheet names/order differ. expected={EXPECTED_SHEETS} actual={wb.sheetnames}")

raw_h, raw = sheet_records(wb["raw_reviews"])
cln_h, cln = sheet_records(wb["cleaned_reviews"])
thm_h, thm = sheet_records(wb["theme_tags"])
src_h, src = sheet_records(wb["source_log"])

THEMES = [h for h in thm_h if h not in
          ("review_id", "platform", "reviewer_name", "review_excerpt", "reason")]

info("STRUCT", f"raw_reviews={len(raw)} cleaned_reviews={len(cln)} "
               f"theme_tags={len(thm)} source_log={len(src)} themes={len(THEMES)}")

# raw_reviews has no id column; theme_tags does. Reconstruct the pairing by
# position, which is how build_sopo_workbook.py writes both sheets.
if len(raw) != len(thm):
    err("S2", f"raw_reviews ({len(raw)}) and theme_tags ({len(thm)}) row counts differ; "
              "positional pairing between the sheets is unreliable.")
ids = [t["review_id"] for t in thm]

# --------------------------------------------------------------------------
# 1. raw vs cleaned
# --------------------------------------------------------------------------
if len(raw) != len(cln):
    warn("C1", f"cleaned_reviews ({len(cln)}) != raw_reviews ({len(raw)}) rows.")
else:
    diffs = [i for i, (a, b) in enumerate(zip(raw, cln)) if a != b]
    if not diffs:
        warn("C2", "cleaned_reviews is a byte-for-byte copy of raw_reviews. No cleaning, "
                   "normalisation, or quarantining was applied, so the sheet adds no "
                   "information over raw_reviews.")

# --------------------------------------------------------------------------
# 2. theme_tags integrity
# --------------------------------------------------------------------------
if len(ids) != len(set(ids)):
    dupes = [k for k, v in Counter(ids).items() if v > 1]
    err("T1", f"Duplicate review_id in theme_tags: {dupes}")

for t in thm:
    for th in THEMES:
        v = t[th]
        if v not in (0, 1):
            err("T2", f"{t['review_id']}: theme '{th}' = {v!r}, expected 0 or 1.")
    if not str(t["reason"]).strip():
        err("T3", f"{t['review_id']}: empty reason.")
    if sum(int(t[th]) for th in THEMES) == 0:
        warn("T4", f"{t['review_id']}: no theme flagged at all.")

# dead theme columns
for th in THEMES:
    total = sum(int(t[th]) for t in thm)
    if total == 0:
        warn("T5", f"Theme column '{th}' is 0 for every one of the {len(thm)} rows — "
                   "carries no signal as collected.")
    elif total == 1:
        warn("T6", f"Theme column '{th}' has a single hit (n=1) — too thin to support "
                   "any conclusion.")

# platform / reviewer_name must agree between raw and theme sheets
for i, (r, t) in enumerate(zip(raw, thm)):
    if r["platform"] != t["platform"]:
        err("T7", f"{t['review_id']}: platform mismatch raw={r['platform']!r} "
                  f"theme_tags={t['platform']!r}")
    if r["reviewer_name"] != t["reviewer_name"]:
        err("T8", f"{t['review_id']}: reviewer_name mismatch raw={r['reviewer_name']!r} "
                  f"theme_tags={t['reviewer_name']!r}")
    ex = str(t["review_excerpt"])
    base = ex[:-3] if ex.endswith("...") else ex
    if base and not str(r["review_text"]).startswith(base):
        err("T9", f"{t['review_id']}: review_excerpt is not a prefix of review_text.")

# --------------------------------------------------------------------------
# 3. Theme flags vs the text that is supposed to justify them
# --------------------------------------------------------------------------
# Keyword evidence per theme. A flag with zero lexical support in review_text
# is not automatically wrong, but it should be traceable to the reason string.
EVIDENCE = {
    "portion_size": r"\b(portion|small|smaller|size|tiny|large|generous|filling|amount)\b",
    "value_for_money": r"\b(price|pricing|prices|priced|value|cheap|affordable|expensive|\$|cost|worth)\b",
    "service_experience": r"\b(staff|service|server|friendly|rude|waiter|employee|cashier)\b",
    "ambience_store_environment": r"\b(seating|seat|space|ambience|ambiance|decor|atmosphere|crowded|clean|table)\b",
    "health_freshness": r"\b(fresh|freshness|healthy|health|clean ingredients|nutriti)\b",
    "catering_group_order": r"\b(cater|catering|group|party|office|platter|bulk)\b",
    "repeat_visit_intent": r"\b(again|come back|go back|return|every single time|regular|keeping this|handy spot)\b",
}
for r, t in zip(raw, thm):
    text = str(r["review_text"]).lower()
    for th, pat in EVIDENCE.items():
        if int(t[th]) == 1 and not re.search(pat, text):
            err("F1", f"{t['review_id']}: '{th}'=1 but review_text contains no supporting "
                      f"term. text={str(r['review_text'])[:110]!r}")

# --------------------------------------------------------------------------
# 4. Provenance rules from the build script's own docstring
# --------------------------------------------------------------------------
PARA_TAG = "[PARAPHRASED SEARCH SUMMARY - NOT VERBATIM]"
for r, t in zip(raw, thm):
    rid = t["review_id"]
    txt = str(r["review_text"])
    is_para = txt.startswith(PARA_TAG)
    if is_para:
        # rule: reviewer_name / date / rating must be blank for paraphrased rows
        for field in ("reviewer_name", "review_date", "rating"):
            if str(r[field]).strip():
                err("P1", f"{rid}: paraphrased row but '{field}' is populated "
                          f"({r[field]!r}) — violates the stated no-guessing rule.")
        if "NOT VERBATIM" not in str(r["notes"]).upper():
            warn("P2", f"{rid}: paraphrased row whose notes do not repeat the "
                       "NOT VERBATIM caveat.")
    if not str(r["source_url"]).strip():
        err("P3", f"{rid}: source_url is empty — row is unciteable.")
    if not str(r["review_url"]).strip():
        warn("P4", f"{rid}: no direct review_url; only a source_url. Claim cannot be "
                   "traced to a single post.")
    if not txt.strip():
        err("P5", f"{rid}: review_text is empty.")

# --------------------------------------------------------------------------
# 5. source_log coverage vs the rows that actually exist
# --------------------------------------------------------------------------
src_blob = " ".join(str(v) for s in src for v in s.values())
cited = set(re.findall(r"\bR\d{2}\b", src_blob))
missing = [i for i in ids if i not in cited]
if missing:
    err("L1", f"Review ids present in the data but never referenced anywhere in "
              f"source_log: {missing}")
ghost = sorted(cited - set(ids))
if ghost:
    err("L2", f"source_log references review ids that do not exist: {ghost}")

# per-platform counts claimed in source_log vs actual rows
plat_actual = defaultdict(list)
for r, t in zip(raw, thm):
    plat_actual[str(r["platform"])].append(t["review_id"])

for s in src:
    label = str(s["source"])
    claim = str(s["reviews_collected"])
    listed = set(re.findall(r"\bR\d{2}\b", claim))
    if not listed:
        continue
    # which platform bucket does this source correspond to?
    for plat, rids in plat_actual.items():
        if not listed & set(rids):
            continue
        extra = set(rids) - listed
        if extra:
            err("L3", f"source_log['{label}'] claims {sorted(listed)} for platform "
                      f"'{plat}', but that platform actually has {sorted(rids)} — "
                      f"missing {sorted(extra)}.")
        m = re.match(r"\s*(\d+)", claim)
        nums = re.findall(r"(\d+)\s*(?:verbatim|captions?|paraphrased)", claim)

# numeric claim vs listed ids inside each source_log cell
for s in src:
    claim = str(s["reviews_collected"])
    listed = re.findall(r"\bR\d{2}\b", claim)
    for n, word in re.findall(r"(\d+)\s+(captions?\s+captured|paraphrased\s+summar\w+)", claim):
        if listed and int(n) != len(listed):
            err("L4", f"source_log['{s['source']}']: says '{n} {word}' but lists "
                      f"{len(listed)} ids {listed}.")

# every source_url used by a row should appear somewhere in source_log
src_urls = {str(s["url"]) for s in src}
for r, t in zip(raw, thm):
    u = str(r["source_url"])
    if not u:
        continue
    host = re.sub(r"^https?://(www\.)?", "", u).split("/")[0]
    if not any(host in su for su in src_urls) and not any(host in str(v) for s in src for v in s.values()):
        warn("L5", f"{t['review_id']}: source host '{host}' has no source_log entry "
                   "documenting how it was accessed.")

# --------------------------------------------------------------------------
# 6. Independently verifiable: TikTok IDs encode the post timestamp
# --------------------------------------------------------------------------
def tiktok_ts(url):
    m = re.search(r"/video/(\d{15,25})", str(url))
    if not m:
        return None
    return dt.datetime.utcfromtimestamp(int(m.group(1)) >> 32)

for r, t in zip(raw, thm):
    ts = tiktok_ts(r["review_url"])
    if ts is None:
        continue
    true_date = ts.strftime("%Y-%m-%d")
    recorded = str(r["review_date"])
    if not recorded or recorded.lower().startswith("unknown"):
        err("D1", f"{t['review_id']}: review_date recorded as {recorded!r}, but the "
                  f"TikTok video id decodes to {true_date} ({ts:%H:%M} UTC). The date "
                  "was recoverable without any network access.")
    else:
        got = re.search(r"(\d{4}-\d{2}-\d{2})", recorded)
        if got:
            delta = abs((dt.date.fromisoformat(got.group(1)) - ts.date()).days)
            if delta > 1:
                err("D2", f"{t['review_id']}: review_date {got.group(1)} is {delta} days "
                          f"from the id-derived post date {true_date}.")
            else:
                info("D3", f"{t['review_id']}: recorded date {got.group(1)} agrees with "
                           f"id-derived {true_date} (UTC/local boundary).")

# --------------------------------------------------------------------------
# 7. Staleness
# --------------------------------------------------------------------------
accessed = {str(s["date_accessed"]) for s in src if re.match(r"\d{4}-\d{2}-\d{2}", str(s["date_accessed"]))}
if accessed:
    newest = max(dt.date.fromisoformat(a) for a in accessed)
    age = (dt.date.today() - newest).days
    if age > 60:
        warn("A1", f"Newest date_accessed is {newest} — {age} days old. Ratings and "
                   "review counts quoted in source_log should be re-checked before use.")

# --------------------------------------------------------------------------
# 8. Aggregate figures quoted in prose
# --------------------------------------------------------------------------
for s in src:
    blob = " ".join(str(v) for v in s.values())
    for m in re.finditer(r"(\d[\d,]*(?:\.\d)?)\s*(?:star|stars|reviews|photos)", blob):
        pass
    if re.search(r"\bNOT\b.{0,20}\bverif|could not be verified|not confirmed", blob, re.I):
        info("Q1", f"source_log['{s['source']}'] carries an explicitly unverified "
                   "aggregate figure — do not quote it as fact.")

# --------------------------------------------------------------------------
# Report
# --------------------------------------------------------------------------
def dump(title, items):
    print(f"\n{title} ({len(items)})")
    print("-" * 72)
    if not items:
        print("  none")
    for code, msg in items:
        print(f"  [{code}] {msg}")


print("=" * 72)
print(f"DATA AUDIT — {PATH}")
print("=" * 72)
dump("ERRORS", ERRORS)
dump("WARNINGS", WARNS)
dump("INFO", INFOS)
print()
print(f"TOTAL: {len(ERRORS)} errors, {len(WARNS)} warnings, {len(INFOS)} info")
sys.exit(1 if ERRORS else 0)
